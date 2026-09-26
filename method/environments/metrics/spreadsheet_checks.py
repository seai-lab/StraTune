"""Canonical task-normalized SpreadsheetBench value scorer."""

from __future__ import annotations

import datetime as dt
import re
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils.cell import column_index_from_string, get_column_letter

CELL_RE = re.compile(r"^\$?([A-Za-z]+)\$?(\d+)$")
COL_RE = re.compile(r"^\$?([A-Za-z]+)$")
ROW_RE = re.compile(r"^\$?(\d+)$")


@dataclass(frozen=True)
class AnswerRange:
    sheet: str
    cells: tuple[str, ...]


@dataclass(frozen=True)
class CaseScore:
    hard_pass: float
    cell_accuracy: float
    correct_cells: int
    total_cells: int
    errors: tuple[str, ...]


@dataclass(frozen=True)
class TaskScore:
    hard_pass: float
    case_pass_fraction: float
    cell_accuracy: float
    case_count: int
    errors: tuple[str, ...]


def transform_value(value: Any) -> Any:
    """Match the official value-only normalization."""

    if isinstance(value, bool):
        return round(float(value), 2)
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    if isinstance(value, dt.time):
        return str(value)[:-3]
    if isinstance(value, dt.datetime):
        origin = dt.datetime(1899, 12, 30)  # noqa: DTZ001 - Excel uses naive dates.
        delta = value - origin
        serial = delta.days + delta.seconds / 86400.0
        return round(serial, 0)
    if isinstance(value, str):
        try:
            return round(float(value), 2)
        except ValueError:
            return value
    return value


def values_equal(left: Any, right: Any) -> bool:
    left = transform_value(left)
    right = transform_value(right)
    if left in (None, "") and right in (None, ""):
        return True
    return type(left) is type(right) and left == right


def split_answer_specs(value: str) -> tuple[str, ...]:
    """Split range specs without breaking commas inside a sheet name.

    SpreadsheetBench contains both correctly quoted sheet names and rows
    written without quoting, or with unmatched or misplaced quotes. A comma can
    delimit a new range only after the current fragment has crossed its
    sheet/range ``!`` boundary.
    """

    if "!" not in value:
        return tuple(part.strip() for part in value.split(",") if part.strip())

    parts: list[str] = []
    current: list[str] = []
    crossed_boundary = False
    for char in value:
        if char == "," and crossed_boundary:
            part = "".join(current).strip()
            if part:
                parts.append(part)
            current = []
            crossed_boundary = False
        else:
            current.append(char)
            crossed_boundary = crossed_boundary or char == "!"
    final = "".join(current).strip()
    if final:
        parts.append(final)
    return tuple(parts)


def _unquote_sheet(value: str) -> str:
    value = value.strip().strip("\u00a0").strip("'\"‘’")
    if len(value) >= 2 and value[0] == value[-1] and value[0] in {"'", '"'}:
        quote = value[0]
        return value[1:-1].replace(quote * 2, quote)
    return value


def _split_sheet(spec: str, default_sheet: str) -> tuple[str, str]:
    if "!" not in spec:
        return default_sheet, spec.strip().strip("'\"")
    sheet, range_text = spec.split("!", 1)
    if "!" in range_text:
        range_text = range_text.rsplit("!", 1)[1]
    return _unquote_sheet(sheet), range_text.strip().strip("'\"")


def _sheet_key(value: str) -> str:
    return "".join(
        char
        for char in value.casefold().strip().strip("\u00a0'\"‘’")
        if not char.isspace()
    )


def _resolve_sheet(workbook: Any, requested: str) -> str:
    if requested in workbook.sheetnames:
        return requested
    key = _sheet_key(requested)
    matches = [name for name in workbook.sheetnames if _sheet_key(name) == key]
    if len(matches) == 1:
        return matches[0]
    if len(workbook.sheetnames) == 1:
        return workbook.sheetnames[0]
    raise KeyError(f"worksheet not found: {requested}")


def _parse_endpoint(value: str) -> tuple[str | None, int | None]:
    value = value.strip()
    if match := CELL_RE.fullmatch(value):
        return match.group(1).upper(), int(match.group(2))
    if match := COL_RE.fullmatch(value):
        return match.group(1).upper(), None
    if match := ROW_RE.fullmatch(value):
        return None, int(match.group(1))
    raise ValueError(f"invalid cell/range endpoint: {value!r}")


def expand_range(range_text: str, *, max_row: int, max_column: int) -> tuple[str, ...]:
    clean = range_text.strip().strip('"')
    if ":" not in clean:
        column, row = _parse_endpoint(clean)
        if column is None or row is None:
            raise ValueError(f"single coordinate must be a cell: {clean!r}")
        return (f"{column}{row}",)

    left_text, right_text = clean.split(":", 1)
    left_column, left_row = _parse_endpoint(left_text)
    right_column, right_row = _parse_endpoint(right_text)

    # Canonicalize a common malformed form such as BD2:308.
    if left_column is not None and left_row is not None and right_column is None:
        right_column = left_column
    if left_column is None and left_row is not None and right_column is not None:
        left_column = right_column

    if left_column is None and right_column is None:
        start_column, end_column = 1, max(1, max_column)
    elif left_column is not None and right_column is not None:
        start_column = column_index_from_string(left_column)
        end_column = column_index_from_string(right_column)
    else:
        raise ValueError(f"incompatible range columns: {clean!r}")

    if left_row is None and right_row is None:
        start_row, end_row = 1, max(1, max_row)
    elif left_row is not None and right_row is not None:
        start_row, end_row = left_row, right_row
    else:
        raise ValueError(f"incompatible range rows: {clean!r}")

    if start_column > end_column or start_row > end_row:
        raise ValueError(f"descending range is unsupported: {clean!r}")
    return tuple(
        f"{get_column_letter(column)}{row}"
        for column in range(start_column, end_column + 1)
        for row in range(start_row, end_row + 1)
    )


def answer_ranges(
    workbook: Any,
    answer_position: str,
    answer_sheet: str = "",
) -> tuple[AnswerRange, ...]:
    default_sheet = answer_sheet.strip() or workbook.sheetnames[0]
    specs = split_answer_specs(answer_position or "")
    sheet_specs: tuple[tuple[str, str], ...]
    if specs and all("!" not in spec for spec in specs) and "," in default_sheet:
        if default_sheet in workbook.sheetnames:
            sheet_specs = tuple((default_sheet, spec) for spec in specs)
        else:
            sheets = tuple(
                _unquote_sheet(part)
                for part in default_sheet.split(",")
                if part.strip()
            )
            if len(specs) == 1:
                sheet_specs = tuple((sheet, specs[0]) for sheet in sheets)
            elif len(specs) == len(sheets):
                sheet_specs = tuple(zip(sheets, specs))
            else:
                raise ValueError(
                    "answer_sheet and answer_position lists have incompatible lengths"
                )
    else:
        sheet_specs = tuple(_split_sheet(spec, default_sheet) for spec in specs)
    ranges = []
    for requested_sheet, range_text in sheet_specs:
        sheet = _resolve_sheet(workbook, requested_sheet)
        worksheet = workbook[sheet]
        ranges.append(
            AnswerRange(
                sheet=sheet,
                cells=expand_range(
                    range_text,
                    max_row=worksheet.max_row,
                    max_column=worksheet.max_column,
                ),
            )
        )
    if not ranges:
        raise ValueError("answer_position contains no scoreable range")
    return tuple(ranges)


def score_case(
    prediction_path: str | Path,
    gold_path: str | Path,
    *,
    answer_position: str,
    answer_sheet: str = "",
) -> CaseScore:
    prediction_path = Path(prediction_path)
    gold_path = Path(gold_path)
    if not prediction_path.exists():
        return CaseScore(0.0, 0.0, 0, 0, ("missing_output",))
    try:
        gold = openpyxl.load_workbook(gold_path, data_only=True)
        prediction = openpyxl.load_workbook(prediction_path, data_only=True)
    except Exception as exc:  # noqa: BLE001
        return CaseScore(0.0, 0.0, 0, 0, (f"load_error:{type(exc).__name__}",))

    correct = 0
    total = 0
    errors: list[str] = []
    try:
        try:
            ranges = answer_ranges(gold, answer_position, answer_sheet)
        except KeyError as exc:
            return CaseScore(0.0, 0.0, 0, 0, (f"metadata_{exc.args[0]}",))
        except ValueError as exc:
            return CaseScore(0.0, 0.0, 0, 0, (f"metadata_error:{exc}",))
        for target in ranges:
            if target.sheet not in prediction.sheetnames:
                errors.append(f"missing_sheet:{target.sheet}")
                total += len(target.cells)
                continue
            gold_sheet = gold[target.sheet]
            prediction_sheet = prediction[target.sheet]
            for coordinate in target.cells:
                total += 1
                if values_equal(
                    gold_sheet[coordinate].value,
                    prediction_sheet[coordinate].value,
                ):
                    correct += 1
                elif len(errors) < 20:
                    errors.append(f"value_mismatch:{target.sheet}!{coordinate}")
    finally:
        gold.close()
        prediction.close()
    accuracy = correct / total if total else 0.0
    return CaseScore(float(total > 0 and correct == total), accuracy, correct, total, tuple(errors))


def summarize_task(cases: Iterable[CaseScore]) -> TaskScore:
    rows = tuple(cases)
    if not rows:
        return TaskScore(0.0, 0.0, 0.0, 0, ("no_cases",))
    errors = tuple(sorted({error for row in rows for error in row.errors}))
    return TaskScore(
        hard_pass=float(all(row.hard_pass == 1.0 for row in rows)),
        case_pass_fraction=sum(row.hard_pass for row in rows) / len(rows),
        cell_accuracy=sum(row.cell_accuracy for row in rows) / len(rows),
        case_count=len(rows),
        errors=errors,
    )
