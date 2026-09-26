"""SpreadsheetBench single-call codegen environment (the single evaluator shared
by training and test).

Prompt construction, sandbox execution and scoring in this file are fixed; the
environment's version string (``VERSION``) is recorded in every result.

Protocol (mirrors SkillOpt's one-code / all-cases semantics)
-----------------------------------------------------------
1. ONE code-generation LLM call per task: system = the supplied
   ``system_prompt``; user = instruction + instruction_type +
   answer_position + structural preview of case 1's input workbook (the
   SkillOpt preview format: per sheet header with dimensions, first 5 rows x
   20 columns as ``coordinate=value`` cells, 40-char cell cap) + the fixed
   task paragraph asking for a single ```python``` block operating on
   ``INPUT_PATH`` -> ``OUTPUT_PATH``.
2. The first fenced code block is extracted (SkillOpt ``extract_code``).
3. For EVERY case of the task the same code is executed in a subprocess
   sandbox against a fresh copy of that case's input workbook inside a
   task-scoped tmp workdir (timeout 120 s per case; Python-level network
   disabled via a socket shim; cwd isolated; INPUT_PATH/OUTPUT_PATH
   assignments in model code stripped and re-injected -- identical to the
   SkillOpt executor template).
4. Each produced workbook is scored against the case's answer/golden
   workbook with the official-semantics scorer
   ``method.environments.metrics.spreadsheet_checks.score_case`` and aggregated with
   ``summarize_task``: primary = ``hard_pass`` (all cases pass), soft =
   ``case_pass_fraction``, plus ``cell_accuracy``.

Both data layouts are handled transparently because the
``method.environments.datasets.spreadsheetbench`` loader normalizes them into case records
with ``input_path``/``answer_path``: the 912 payload (``*_input.xlsx`` /
``*_answer.xlsx``, up to 3 cases) and the verified_400 payload
(``*_init.xlsx`` / ``*_golden.xlsx`` or bare ``initial.xlsx`` /
``golden.xlsx``, 1 case).

Deviations from the SkillOpt defaults: single-call mode only (no
multi-turn error feedback), a global preview character cap (20 000), the
sandbox socket shim + isolated cwd, no post-call rate-limit sleep, and
scoring via ``method.environments.metrics.spreadsheet_checks`` instead of
SkillOpt's evaluator (same official value-comparison semantics, more
workbook-quirk handling).
"""

from __future__ import annotations

import os
import re
import shutil
import subprocess
import sys
import textwrap
from pathlib import Path
from typing import Any

import openpyxl

from method.environments.bedrock_client import TARGET_MODEL
from method.environments.bedrock_client import BedrockConverseClient
from method.environments.metrics.spreadsheet_checks import CaseScore
from method.environments.metrics.spreadsheet_checks import score_case
from method.environments.metrics.spreadsheet_checks import summarize_task

VERSION = "1.0.0"

PREVIEW_MAX_ROWS = 5
PREVIEW_MAX_COLS = 20
PREVIEW_CELL_CHARS = 40
PREVIEW_MAX_CHARS = 20_000

EXEC_TIMEOUT_SECONDS = 120

# ---------------------------------------------------------------------------
# Prompt construction (SkillOpt-faithful)
# ---------------------------------------------------------------------------


def preview_workbook(
    path: str | Path,
    max_rows: int = PREVIEW_MAX_ROWS,
    max_cols: int = PREVIEW_MAX_COLS,
    max_chars: int = PREVIEW_MAX_CHARS,
) -> str:
    """Structural preview of a workbook (SkillOpt ``_preview_workbook``).

    Per sheet: a header line with name, dimensions, max_row/max_col, then the
    first ``max_rows`` rows x ``max_cols`` columns as ``coordinate=value``
    cells (values truncated to 40 chars). A global character cap
    (``max_chars``) is applied on top of the reference format.
    """
    wb = openpyxl.load_workbook(str(path), data_only=False)
    chunks: list[str] = []
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            chunks.append(
                f"## Sheet: {sheet_name}  "
                f"(dim={ws.dimensions}, max_row={ws.max_row}, max_col={ws.max_column})"
            )
            for row in ws.iter_rows(
                min_row=1,
                max_row=min(ws.max_row, max_rows),
                max_col=min(ws.max_column, max_cols),
                values_only=False,
            ):
                cells = []
                for cell in row:
                    v = cell.value
                    if v is None:
                        cells.append(f"{cell.coordinate}=")
                    else:
                        s = str(v)
                        if len(s) > PREVIEW_CELL_CHARS:
                            s = s[: PREVIEW_CELL_CHARS - 3] + "..."
                        cells.append(f"{cell.coordinate}={s}")
                chunks.append(" | ".join(cells))
            if ws.max_row > max_rows:
                chunks.append(f"... ({ws.max_row - max_rows} more rows)")
            chunks.append("")
    finally:
        wb.close()
    preview = "\n".join(chunks)
    if len(preview) > max_chars:
        preview = preview[:max_chars] + "\n... (preview truncated)"
    return preview


def extract_code(text: str) -> str:
    """Extract the first fenced code block (SkillOpt ``extract_code``)."""
    if "```" not in text:
        return text.strip()
    start = text.find("```")
    nl = text.find("\n", start)
    end = text.find("```", nl + 1)
    if nl == -1 or end == -1:
        return text.strip()
    return text[nl + 1 : end].strip()


def build_user_prompt(
    instruction: str,
    preview: str,
    instruction_type: str = "",
    answer_position: str = "",
) -> str:
    """User prompt, same structure as SkillOpt ``_build_user``
    (non-diagnostic path)."""
    extra = ""
    if instruction_type:
        extra += f"\nInstruction type: {instruction_type}"
    if answer_position:
        extra += f"\nExpected answer position: {answer_position}"
    return (
        f"# Instruction\n{instruction}\n{extra}\n\n"
        f"# Input spreadsheet preview\n{preview}\n\n"
        "# Task\n"
        "Write a Python script that reads the workbook from the variable `INPUT_PATH`, "
        "applies the instruction, and writes the modified workbook to `OUTPUT_PATH`. "
        "Preserve all other cells unchanged. "
        "The preview may be truncated — do not hardcode row counts or assume the data ends at the last previewed row; "
        "iterate over all actual rows in the workbook instead. "
        "Return only a ```python``` code block."
    )


# ---------------------------------------------------------------------------
# Sandbox execution (SkillOpt executor template + hardening)
# ---------------------------------------------------------------------------

_PATH_ASSIGN_RE = re.compile(r"^\s*(INPUT_PATH|OUTPUT_PATH)\s*=\s*.+$", re.MULTILINE)

_RUNNER_TEMPLATE = textwrap.dedent(
    """
    import os, sys, traceback
    # -- sandbox: disable Python-level network access (best effort) --------
    import socket as _socket
    def _no_network(*args, **kwargs):
        raise OSError("network access is disabled in the codegen sandbox")
    _socket.socket = _no_network
    _socket.create_connection = _no_network
    _socket.getaddrinfo = _no_network
    INPUT_PATH = {input_path!r}
    OUTPUT_PATH = {output_path!r}
    try:
    {user_code_indented}
    except Exception:
        traceback.print_exc()
        sys.exit(2)
    """
)


def _strip_path_assignments(code: str) -> str:
    return _PATH_ASSIGN_RE.sub("", code)


def run_code_sandboxed(
    code: str,
    input_path: str | Path,
    output_path: str | Path,
    workdir: str | Path,
    timeout: int = EXEC_TIMEOUT_SECONDS,
) -> tuple[bool, str]:
    """Execute generated code in a subprocess sandbox.

    Wraps the code in the SkillOpt runner template (INPUT_PATH/OUTPUT_PATH
    injected, model-written assignments stripped) plus a socket shim, runs it
    with cwd=``workdir`` and proxy env vars removed, and returns
    ``(ok, error_text)``. ``ok`` requires exit code 0 AND the output file to
    exist. Timeout -> ``(False, "timeout after {timeout}s")``.
    """
    workdir = Path(workdir)
    workdir.mkdir(parents=True, exist_ok=True)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    cleaned = _strip_path_assignments(code)
    indented = textwrap.indent(cleaned, "    ")
    script = _RUNNER_TEMPLATE.format(
        input_path=str(input_path),
        output_path=str(output_path),
        user_code_indented=indented,
    )
    script_path = workdir / "_runner.py"
    script_path.write_text(script, encoding="utf-8")
    env = {
        key: value
        for key, value in os.environ.items()
        if key
        not in {
            "HTTP_PROXY", "HTTPS_PROXY", "http_proxy", "https_proxy",
            "AWS_ACCESS_KEY_ID", "AWS_SECRET_ACCESS_KEY", "AWS_SESSION_TOKEN",
        }
    }
    try:
        proc = subprocess.run(
            [sys.executable, str(script_path)],
            capture_output=True,
            text=True,
            timeout=timeout if timeout and timeout > 0 else None,
            cwd=str(workdir),
            env=env,
        )
    except subprocess.TimeoutExpired:
        return False, f"timeout after {timeout}s"
    if proc.returncode != 0:
        return False, (proc.stdout + "\n" + proc.stderr).strip()
    if not output_path.exists():
        return False, "output file was not created"
    return True, ""


# ---------------------------------------------------------------------------
# Environment
# ---------------------------------------------------------------------------


def _task_fields(task: Any) -> dict:
    """Normalize a SpreadsheetBenchTask or mapping into plain fields."""
    if isinstance(task, dict):
        return {
            "task_id": str(task["task_id"]),
            "instruction": str(task["instruction"]),
            "instruction_type": str(task.get("instruction_type") or ""),
            "answer_position": str(task.get("answer_position") or ""),
            "answer_sheet": str(task.get("answer_sheet") or ""),
            "cases": [
                {
                    "case": c["case"],
                    "input_path": Path(c["input_path"]),
                    "answer_path": Path(c["answer_path"]),
                }
                for c in task["cases"]
            ],
        }
    return {
        "task_id": str(task.task_id),
        "instruction": str(task.instruction),
        "instruction_type": str(task.instruction_type or ""),
        "answer_position": str(task.answer_position or ""),
        "answer_sheet": str(task.answer_sheet or ""),
        "cases": [
            {
                "case": c.case,
                "input_path": Path(c.input_path),
                "answer_path": Path(c.answer_path),
            }
            for c in task.cases
        ],
    }


class SpreadsheetCodegenEnv:
    """Single-call codegen host over SpreadsheetBench tasks.

    ``evaluate_task(task, system_prompt, client)`` makes exactly one code
    generation call (preview of case 1), applies the code to all cases in a
    task-scoped tmp workdir, scores with
    ``method.environments.metrics.spreadsheet_checks``, and returns the task
    record (primary metric: ``hard_pass``).
    """

    def __init__(
        self,
        work_root: str | Path,
        model: str = TARGET_MODEL,
        max_tokens: int = 16_384,
        seed: int = 42,
        exec_timeout: int = EXEC_TIMEOUT_SECONDS,
    ) -> None:
        self.work_root = Path(work_root)
        self.work_root.mkdir(parents=True, exist_ok=True)
        self.model = model
        self.max_tokens = max_tokens
        self.seed = seed
        self.exec_timeout = exec_timeout

    # -- code application (no LLM) ---------------------------------------
    def apply_code_to_task(self, code: str, fields: dict, workdir: Path) -> dict:
        """Run ``code`` on every case and score. Pure execution + scoring."""
        case_results: list[dict] = []
        case_scores: list[CaseScore] = []
        for case in fields["cases"]:
            case_no = str(case["case"])
            case_dir = workdir / f"case_{case_no}"
            case_dir.mkdir(parents=True, exist_ok=True)
            local_input = case_dir / f"input_{case['input_path'].name}"
            shutil.copy2(case["input_path"], local_input)
            output_path = case_dir / "output.xlsx"
            if code.strip():
                ok, error = run_code_sandboxed(
                    code,
                    local_input,
                    output_path,
                    case_dir,
                    timeout=self.exec_timeout,
                )
            else:
                ok, error = False, "no python code block extracted"
            score = score_case(
                output_path,
                case["answer_path"],
                answer_position=fields["answer_position"],
                answer_sheet=fields["answer_sheet"],
            )
            case_scores.append(score)
            case_results.append(
                {
                    "case": case_no,
                    "exec_ok": ok,
                    "exec_error": error[:2000],
                    "hard_pass": score.hard_pass,
                    "cell_accuracy": score.cell_accuracy,
                    "correct_cells": score.correct_cells,
                    "total_cells": score.total_cells,
                    "score_errors": list(score.errors)[:20],
                }
            )
        task_score = summarize_task(case_scores)
        return {
            "hard_pass": task_score.hard_pass,
            "case_pass_fraction": task_score.case_pass_fraction,
            "cell_accuracy": task_score.cell_accuracy,
            "case_count": task_score.case_count,
            "errors": list(task_score.errors),
            "case_results": case_results,
        }

    # -- full protocol -----------------------------------------------------
    def evaluate_task(
        self,
        task: Any,
        system_prompt: str,
        client: BedrockConverseClient,
    ) -> dict:
        fields = _task_fields(task)
        task_id = fields["task_id"]
        if not fields["cases"]:
            raise ValueError(f"task {task_id} has no cases")

        preview_input = fields["cases"][0]["input_path"]
        try:
            preview = preview_workbook(preview_input)
        except Exception as exc:  # noqa: BLE001 - reference behavior
            preview = f"(failed to preview workbook: {exc})"
        user_prompt = build_user_prompt(
            fields["instruction"],
            preview,
            instruction_type=fields["instruction_type"],
            answer_position=fields["answer_position"],
        )
        response = client.converse(
            self.model,
            [{"role": "user", "content": [{"text": user_prompt}]}],
            system=system_prompt,
            max_tokens=self.max_tokens,
            temperature=0.0,
            seed=self.seed,
            tags={
                "env": "spreadsheet_codegen",
                "env_version": VERSION,
                "task_id": task_id,
            },
        )
        code = extract_code(response["text"])
        workdir = self.work_root / task_id
        applied = self.apply_code_to_task(code, fields, workdir)
        usage = response.get("usage", {})
        return {
            "task_id": task_id,
            "primary": applied["hard_pass"],
            **applied,
            "code": code,
            "response_text": response["text"],
            "stop_reason": response.get("stop_reason"),
            "cached": response.get("cached", False),
            "usage": {
                "input_tokens": int(usage.get("inputTokens", 0)),
                "output_tokens": int(usage.get("outputTokens", 0)),
            },
            "model": self.model,
            "env_version": VERSION,
        }
